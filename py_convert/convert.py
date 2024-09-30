import openpyxl
import docx

def replace_test_results_in_template(xlsx_file_path, template_path, output_path):
    try:
        # Read data from the Excel file
        test_results = {}
        workbook = openpyxl.load_workbook(xlsx_file_path)
        sheet = workbook.active  # Use the active sheet

        # Assuming the first row contains headers
        for row in sheet.iter_rows(min_row=2, values_only=True):  # Skip header row
            if len(row) < 5:  # Ensure there are enough columns
                continue
            
            test_name = row[0]  # Test Name is in the first column
            result = row[1]      # Result is in the second column
            unit = row[2]        # Unit is in the third column
            flag = row[3]        # Flag is in the fourth column (optional)
            reference_interval = row[4]  # Reference Interval in the fifth column (optional)

            # Store the results with units, flags, and reference intervals
            test_results[test_name] = {
                "result": result,
                "unit": unit,
                "flag": flag,
                "reference_interval": reference_interval
            }

        # Debugging output to verify test results
        print("Test Results Loaded:")
        for test_name, details in test_results.items():
            print(f"{test_name}: {details}")

        # Open the Word template
        doc = docx.Document(template_path)

        # Iterate over the tables in the Word document
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for test_name, details in test_results.items():
                        # Create placeholders
                        result_placeholder = f"{{{test_name.lower().replace(' ', '_')}_result}}"
                        unit_placeholder = f"{{{test_name.lower().replace(' ', '_')}_units}}"
                        flag_placeholder = f"{{{test_name.lower().replace(' ', '_')}_flag}}"
                        reference_placeholder = f"{{{test_name.lower().replace(' ', '_')}_reference_interval}}"

                        # Replace placeholders in the cell text
                        if result_placeholder in cell.text:
                            print(f"Replacing '{result_placeholder}' with '{details['result']}' in cell: {cell.text}")
                            cell.text = cell.text.replace(result_placeholder, str(details['result']))

                        if unit_placeholder in cell.text:
                            print(f"Replacing '{unit_placeholder}' with '{details['unit']}' in cell: {cell.text}")
                            cell.text = cell.text.replace(unit_placeholder, str(details['unit']))

                        if flag_placeholder in cell.text:
                            print(f"Replacing '{flag_placeholder}' with '{details['flag']}' in cell: {cell.text}")
                            cell.text = cell.text.replace(flag_placeholder, str(details['flag']))

                        if reference_placeholder in cell.text:
                            print(f"Replacing '{reference_placeholder}' with '{details['reference_interval']}' in cell: {cell.text}")
                            cell.text = cell.text.replace(reference_placeholder, str(details['reference_interval']))

        # Save the updated document
        doc.save(output_path)
        print(f"Updated template generated successfully: {output_path}")

    except Exception as e:
        print(f"An error occurred: {e}")

# File paths
xlsx_file_path = 'Kidney Test.xlsx'  # Path to your XLSX file
template_path = 'Output_file_with_test_results.docx'  # Path to your Word template
output_path = 'Updated_Output_file_with_test_results.docx'  # Path to save the updated Word document

# Call the function to replace test results in the Word template
replace_test_results_in_template(xlsx_file_path, template_path, output_path)
