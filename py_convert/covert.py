import json
from docx import Document

def fill_personal_info_in_template(json_file_path, template_path, output_path):
    try:
        # Load JSON data from file
        with open(json_file_path, 'r', encoding='utf-8') as f:
            json_data = json.load(f)

        # Open the Word template
        doc = Document(template_path)

        # Personal Info and Sample Details fields
        fields = {
            '{patient_code}': json_data.get('patient_code', 'N/A'),
            '{sample_no}': json_data.get('sample_no', 'N/A'),
            '{patient_name}': json_data.get('patient_name', 'N/A'),
            '{age_gender}': json_data.get('age_gender', 'N/A'),
            '{referred_by}': json_data.get('referred_by', 'N/A'),
            '{bill_date}': json_data.get('bill_date', 'N/A'),
            '{sample_collected}': json_data.get('sample_collected', 'N/A'),
            '{sample_received}': json_data.get('sample_received', 'N/A'),
            '{report_completed}': json_data.get('report_completed', 'N/A'),
            '{report_authorized}': json_data.get('report_authorized', 'N/A')
        }

        print("Fields to replace:", fields)  # Debugging output

        # Check and replace placeholders in paragraphs
        for paragraph in doc.paragraphs:
            for placeholder, value in fields.items():
                if placeholder in paragraph.text:
                    print(f"Replacing '{placeholder}' with '{value}' in paragraph: {paragraph.text}")  # Debugging output
                    paragraph.text = paragraph.text.replace(placeholder, value)

        # Save the updated document
        doc.save(output_path)
        print(f"Report generated successfully: {output_path}")

    except Exception as e:
        print(f"An error occurred: {e}")

# File paths
json_file_path = r'C:\Users\isher\OneDrive\Desktop\py covert\data.json'  # Path to your JSON file
template_path = r'C:\Users\isher\OneDrive\Desktop\py covert\template.docx'  # Path to your Word template
output_path = r'C:\Users\isher\OneDrive\Desktop\py covert\Output_file.docx'  # Path to save the filled Word document

# Call the function to fill the Word template
fill_personal_info_in_template(json_file_path, template_path, output_path)
